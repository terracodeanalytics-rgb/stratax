"""Streaming interchange readers. Originals and original units remain authoritative."""
import csv
import gzip
import hashlib
import itertools
import json
import math
import re
from pathlib import Path

PARSER_VERSION = "stratax-import/1.0"
PREVIEW_ROWS = 60
CHUNK_ROWS = 5000
GIB = 1024 ** 3
_output_bytes = 0
_cancel_check = None


def set_cancel_check(callback):
    global _cancel_check
    _cancel_check = callback


class ImportProblem(ValueError):
    def __init__(self, code, message):
        self.code = code
        super().__init__(message)


def scalar(value):
    if value is None:
        return None
    if hasattr(value, "tolist"):
        return scalar(value.tolist())
    if isinstance(value, (tuple, list)):
        return [scalar(v) for v in value]
    if isinstance(value, dict):
        return {str(k): scalar(v) for k, v in value.items()}
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def number(value, null=None):
    if value is None or value == "":
        return None
    try:
        result = float(str(value).replace("D", "E").replace("d", "e"))
        return None if not math.isfinite(result) or (null is not None and result == null) else result
    except (ValueError, TypeError):
        return str(value)


def column(name, unit="", index=0, description=""):
    return {"id": f"column-{index}", "mnemonic": str(name), "unit": str(unit or ""), "description": description}


def issue(code, message, severity="warning"):
    return {"code": code, "message": message, "severity": severity}


def detect(path, hint=""):
    path = Path(path)
    with path.open("rb") as source:
        head = source.read(16384)
    text = head.decode("utf-8-sig", errors="replace")
    suffix = path.suffix.lower()
    if head.startswith(b"LASF"):
        raise ImportProblem("lidar_las", "This is LiDAR LAS, not a well-log LAS file. Export a supported spatial layer.")
    if re.search(r"(?im)^\s*~(?:V|VERSION)", text) and re.search(r"(?im)^\s*VERS\s*\.", text):
        return "las"
    if head.startswith(b"\x89HDF\r\n\x1a\n"):
        return "hdf5"
    if head.startswith(b"SQLite format 3"):
        return "gpkg"
    if head[:4] in (b"II*\x00", b"MM\x00*"):
        return "geotiff" if hint == "spatial" or suffix in (".geotiff",) else "image"
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith((b"\x89PNG", b"\xff\xd8", b"GIF8", b"BM")):
        return "image"
    if head.startswith(b"PK\x03\x04"):
        return {".xlsx": "xlsx", ".docx": "docx", ".epc": "resqml", ".kmz": "kml", ".zip": "bundle"}.get(suffix, "bundle")
    if b"RECORD" in head[:80] and b"V1." in head[:80] or suffix in (".dlis", ".dls"):
        return "dlis"
    if suffix in (".lis", ".lti", ".tap"):
        return "lis"
    if suffix in (".sgy", ".segy", ".segy2"):
        return "segy"
    if suffix == ".su":
        return "su"
    if suffix in (".zgy", ".sgz", ".vds"):
        return suffix[1:]
    if suffix in (".egrid", ".init", ".unrst", ".grid", ".unfmt", ".smspec", ".unsmry"):
        return "eclipse"
    if suffix in (".grdecl", ".data") and re.search(r"\b(SPECGRID|COORD|ZCORN|RUNSPEC|PORO)\b", text, re.I):
        return "grdecl"
    if suffix in (".gri", ".irap", ".zmap", ".cps", ".grd", ".gxf") or text.lstrip().startswith(("-996", "@GRID", "FSASCI", "!")):
        return "surface"
    if re.search(r"(?im)^GOCAD\s", text):
        return "gocad"
    if suffix in (".shp", ".shx", ".dbf", ".prj"):
        return "shapefile"
    if text.lstrip().startswith(("{", "[")):
        return "geojson" if '"Feature' in text or '"coordinates"' in text else "json"
    if text.lstrip().startswith("<"):
        return "kml" if re.search(r"<(?:\w+:)?kml\b", text, re.I) else "xml"
    if suffix == ".xls" or head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if b"\x00" in head:
        raise ImportProblem("unknown_binary", "This binary format is not recognized. Retain the original and export a documented interchange format.")
    return "table"


def digest(path):
    result = hashlib.sha256()
    with open(path, "rb") as source:
        for block in iter(lambda: source.read(4 * 1024 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def output_dataset(output, name, columns, rows, metadata=None, issues=None):
    """Write independently readable compressed chunks without retaining a full table."""
    global _output_bytes
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    dataset_id = "dataset-" + hashlib.sha256((name + str(len(list(output.glob('dataset-*.json'))))).encode()).hexdigest()[:20]
    chunks, preview, count = [], [], 0
    handle, chunk_count, chunk_size, preview_size = None, 0, 0, 0
    minimum, maximum, previous, repeated, reverse = None, None, None, 0, 0
    statistics = [{"min": None, "max": None, "nulls": 0, "repeated": 0, "decreasing": 0} for _ in columns]
    last_values = [None] * len(columns)
    for row in rows:
        row = scalar(row)
        if not isinstance(row, list) or len(row) != len(columns):
            raise ImportProblem("row_width", "Some records do not match the selected column definition. Review the mapping or export settings.")
        for ci, value in enumerate(row):
            stats = statistics[ci]
            if value is None:
                stats['nulls'] += 1
            elif isinstance(value, (int, float)):
                stats['min'] = value if stats['min'] is None else min(value, stats['min'])
                stats['max'] = value if stats['max'] is None else max(value, stats['max'])
                if last_values[ci] is not None:
                    stats['repeated'] += value == last_values[ci]
                    stats['decreasing'] += value < last_values[ci]
                last_values[ci] = value
        encoded = json.dumps(row, allow_nan=False, separators=(",", ":")) + "\n"
        import os
        _output_bytes += len(encoded.encode('utf-8'))
        if _output_bytes > int(os.environ.get('IMPORT_ARTIFACT_BUDGET', str(1024**3))):
            if handle:
                handle.close()
            raise ImportProblem('artifact_budget', 'Parsed data exceeds this job’s artifact budget. Retain the original and export selected channels or smaller source groups.')
        if count % CHUNK_ROWS == 0 and _cancel_check and _cancel_check():
            if handle:
                handle.close()
            raise ImportProblem('cancelled', 'Import cancelled. Retained originals are unchanged.')
        if len(preview) < PREVIEW_ROWS and preview_size + len(encoded) < 250000:
            preview.append(row)
            preview_size += len(encoded)
        index = row[0] if row else None
        if isinstance(index, (float, int)):
            minimum = index if minimum is None else min(index, minimum)
            maximum = index if maximum is None else max(index, maximum)
            if previous is not None:
                repeated += index == previous
                reverse += index < previous
            previous = index
        count += 1
        if handle is None:
            name_part = f"{dataset_id}-{len(chunks):06}.ndjson.gz"
            handle = gzip.open(output / name_part, "wt", encoding="utf-8")
            chunk_count, chunk_size = 0, 0
        handle.write(encoded)
        chunk_count += 1
        chunk_size += len(encoded)
        if chunk_count >= CHUNK_ROWS or chunk_size >= 4 * 1024 * 1024:
            handle.close()
            chunks.append({"name": name_part, "rows": chunk_count, "size": (output / name_part).stat().st_size})
            handle = None
    if handle:
        handle.close()
        chunks.append({"name": name_part, "rows": chunk_count, "size": (output / name_part).stat().st_size})
    notes = list(issues or [])
    if repeated:
        notes.append(issue("repeated_index", f"{repeated} adjacent repeated index values retained."))
    if reverse:
        notes.append(issue("decreasing_index", f"{reverse} decreasing index steps retained; confirm sampling direction."))
    result = {"schemaVersion": 1, "id": dataset_id, "name": name, "columns": columns, "rowCount": count, "preview": preview, "chunks": chunks, "metadata": metadata or {}, "issues": notes, "indexRange": [minimum, maximum], "columnStatistics": statistics, "parserVersion": PARSER_VERSION}
    manifest = json.dumps(result, allow_nan=False)
    if len(manifest) > 900000:
        result["preview"] = []
        result["issues"].append(issue("large_manifest", "Use the chunk index for this large dataset; the inline preview is omitted."))
    (output / (dataset_id + ".json")).write_text(json.dumps(result, allow_nan=False), encoding="utf-8")
    return result


def las_header(line):
    match = re.match(r"\s*([^\s.]+)\s*\.([^\s:]*)\s*([^:]*)(?::(.*))?$", line)
    return {"mnemonic": match[1], "unit": match[2], "value": match[3].strip(), "description": (match[4] or "").strip()} if match else None


def read_las(path, output, options):
    # The section reader handles LAS 3 repeated definition/data sections explicitly.
    headers, curves, datasets, sections = {}, [], [], []
    null, version, section = -999.25, None, ""
    with open(path, encoding=options.get("encoding", "utf-8-sig"), errors="replace") as source:
        while True:
            line = source.readline()
            if not line:
                break
            if line.lstrip().startswith("#") or not line.strip():
                continue
            if line.lstrip().startswith("~"):
                section = line.strip()[1:].split(":")[0].strip()
                upper = section.upper()
                if upper.startswith("C") or upper.endswith("_DEFINITION"):
                    curves = []
                if upper.startswith("A") or upper.endswith("_DATA"):
                    if not curves:
                        raise ImportProblem("missing_curves", "The LAS data section has no matching curve definition.")
                    if version not in ("1.2", "2.0", "3.0"):
                        raise ImportProblem("las_version", "This LAS version is not supported. Export LAS 1.2, 2.0 or 3.0.")
                    pending = []
                    wrapped = headers.get("WRAP", {}).get("value", "NO").upper() == "YES"
                    position = [None]

                    def rows():
                        while True:
                            pos = source.tell()
                            row_line = source.readline()
                            if not row_line:
                                return
                            if row_line.lstrip().startswith("~"):
                                position[0] = pos
                                return
                            if not row_line.strip() or row_line.lstrip().startswith("#"):
                                continue
                            values = re.split(r"\s+|,", row_line.strip())
                            values = [v for v in values if v != ""]
                            if wrapped:
                                pending.extend(values)
                                while len(pending) >= len(curves):
                                    yield [number(v, null) for v in pending[:len(curves)]]
                                    del pending[:len(curves)]
                            else:
                                yield [number(v, null) for v in values]

                    metadata = {"format": "LAS", "version": version, "headers": headers.copy(), "well": headers.get("WELL", {}).get("value", ""), "null": null, "section": section, "coordinatesConfirmed": False}
                    notes = [issue("unknown_units", "Some source curve units are unspecified. Confirm units before scientific use.")] if any(not c["unit"] or c["unit"].upper() == "UNKNOWN" for c in curves) else []
                    if "LAT" in headers or "XCOORD" in headers:
                        notes.append(issue("unconfirmed_crs", "Header coordinates are retained; their source datum must be confirmed before map use."))
                    datasets.append(output_dataset(output, Path(path).name + (" / " + section if datasets else ""), curves.copy(), rows(), metadata, notes))
                    if pending:
                        raise ImportProblem("wrapped_tail", "The final wrapped LAS record is incomplete.")
                    if position[0] is not None:
                        source.seek(position[0])
                continue
            entry = las_header(line)
            if not entry:
                continue
            if section.upper().startswith("C") or section.upper().endswith("_DEFINITION"):
                curves.append(column(entry["mnemonic"], entry["unit"], len(curves), entry["description"]))
            else:
                headers[entry["mnemonic"].upper()] = entry
                if entry["mnemonic"].upper() == "VERS":
                    try:
                        version = f"{float(entry['value']):.1f}"
                    except ValueError:
                        version = None
                if entry["mnemonic"].upper() == "NULL":
                    null = number(entry["value"])
    if not datasets:
        raise ImportProblem("no_data", "No readable LAS data section was found.")
    return datasets


def read_table(path, output, options):
    kind = detect(path)
    columns, iterator = [], None
    skip = int(options.get("headerRow", 1)) - 1
    if kind == "xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        sheet = book[options.get("sheet") or book.sheetnames[0]]
        iterator = sheet.iter_rows(min_row=skip + 1, values_only=True)
    elif kind == "xls":
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        sheet = book.sheet_by_name(options["sheet"]) if options.get("sheet") else book.sheet_by_index(0)
        iterator = (sheet.row_values(i) for i in range(skip, sheet.nrows))
    else:
        source = open(path, encoding=options.get("encoding", "utf-8-sig"), errors="replace", newline="")
        for _ in range(skip):
            next(source, None)
        if options.get("widths"):
            widths = [int(w) for w in options["widths"]]
            if any(w < 1 or w > 10000 for w in widths):
                raise ImportProblem("fixed_widths", "Enter valid fixed-width column sizes.")
            offsets = [0, *itertools.accumulate(widths)]
            iterator = ([line[offsets[i]:offsets[i+1]].strip() for i in range(len(widths))] for line in source)
        else:
            sample = source.read(8192)
            source.seek(0)
            for _ in range(skip):
                next(source, None)
            delimiter = options.get("delimiter")
            if delimiter == "whitespace":
                iterator = (re.split(r"\s+", line.strip()) for line in source if line.strip() and not line.lstrip().startswith(("#", "!")))
            else:
                if not delimiter:
                    try:
                        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                    except csv.Error:
                        delimiter = "whitespace"
                iterator = (re.split(r"\s+", line.strip()) for line in source if line.strip()) if delimiter == "whitespace" else csv.reader(source, delimiter=delimiter)
    first = next(iterator, None)
    if not first:
        raise ImportProblem("empty_table", "The selected sheet or table is empty.")
    if options.get("noHeader"):
        names = options.get("columnNames") or [f"Column {i+1}" for i in range(len(first))]
        iterator = itertools.chain([first], iterator)
    else:
        names = first
    columns = [column(str(name or f"Column {i+1}"), (options.get("units") or {}).get(str(i), ""), i) for i, name in enumerate(names)]
    null = number(options.get("null", ""))
    def rows():
        for row in iterator:
            if not any(v is not None and str(v).strip() for v in row):
                continue
            yield [number(v, null) for v in row]
    try:
        return [output_dataset(output, Path(path).name, columns, rows(), {"format": kind, "mapping": options.get("mapping", {}), "sheet": options.get("sheet", ""), "well": options.get("well", ""), "references": options.get("references", {})})]
    finally:
        if kind == "xlsx":
            book.close()
        elif kind == "xls":
            book.release_resources()
        else:
            source.close()


def validate_mapping(dataset, family, options):
    notes = list(dataset.get("issues", []))
    references = options.get("references", {})
    mapping = options.get("mapping", {})
    if references.get('crs'):
        try:
            from pyproj import CRS
            crs=CRS.from_user_input(str(references['crs'])[:4000])
            dataset['confirmedSourceCRS']=crs.to_string()
            notes=[i for i in notes if i['code']!='unconfirmed_crs']
        except Exception:
            notes.append(issue('invalid_crs','The stated source CRS could not be validated. Correct it or leave it unconfirmed.','blocking'))
    for role, value in mapping.items():
        if value != "":
            try:
                if not 0 <= int(value) < len(dataset['columns']):
                    raise ValueError()
            except (ValueError, TypeError):
                notes.append(issue('invalid_mapping', f'The {role} column is not present in this dataset.', 'blocking'))
    for index, col in enumerate(dataset['columns']):
        effective = options.get('units', {}).get(str(index), '')
        if effective:
            col['effectiveUnit'] = str(effective)[:80]
    if all(c.get('effectiveUnit') or c.get('unit') and c['unit'].upper() != 'UNKNOWN' for c in dataset['columns']):
        notes = [i for i in notes if i['code'] != 'unknown_units']
    required = {"deviation": ["md", "inclination", "azimuth"], "tops": ["name", "depth"], "checkshots": ["depth", "time"], "well_headers": ["well"], "production": ["date"], "spatial": []}.get(family, [])
    if family == "deviation" and options.get("surveyMethod") == "coordinates":
        required = ["md", "x", "y", "z"]
    for field in required:
        if str(mapping.get(field, "")) == "":
            notes.append(issue("mapping_" + field, f"Map the {field} column before activating this dataset.", "blocking"))
    if family in ("logs", "deviation", "tops", "checkshots", "core", "pressure"):
        if not references.get("depthUnit"):
            notes.append(issue("depth_units", "Confirm the source depth units before scientific use."))
        if not references.get("depthReference"):
            notes.append(issue("depth_reference", "Confirm MD, TVD or TVDSS before scientific use."))
    if family == "deviation" and not references.get("northReference"):
        notes.append(issue("north_reference", "North reference is unconfirmed; no trajectory has been calculated."))
    if family == "checkshots" and not references.get("timeConvention"):
        notes.append(issue("time_convention", "Confirm one-way or two-way time and its units."))
    if family in ("seismic", "surfaces", "models", "spatial", "well_headers") and not references.get("crs"):
        notes.append(issue("unconfirmed_crs", "Source CRS is unconfirmed. Geometry is excluded from confirmed map layers."))
    if family in ("seismic", "surfaces", "models") and not references.get("domain"):
        notes.append(issue("domain", "Confirm whether the vertical domain is time or depth."))
    for role, limits in (("inclination", (0, 180)), ("azimuth", (0, 360)), ("latitude", (-90, 90)), ("longitude", (-180, 180))):
        if str(mapping.get(role, "")) != "":
            try:
                index = int(mapping[role])
                if role in ('inclination', 'azimuth') and references.get('angleUnit') == 'radians':
                    limits = (0, math.pi if role == 'inclination' else 2 * math.pi)
                stats = dataset.get('columnStatistics', [{}] * len(dataset['columns']))[index]
                outside = stats.get('min') is not None and (stats['min'] < limits[0] or stats['max'] > limits[1])
                if outside or any(isinstance(row[index], (int, float)) and not limits[0] <= row[index] <= limits[1] for row in dataset.get("preview", [])):
                    notes.append(issue("range_" + role, f"Some {role} values are outside the expected range.", "blocking"))
            except (ValueError, IndexError, TypeError):
                notes.append(issue("invalid_mapping", "A mapped column is not present in this dataset.", "blocking"))
    for role in ('md', 'depth', 'time'):
        if str(mapping.get(role, '')) != '':
            try:
                stats = dataset.get('columnStatistics', [])[int(mapping[role])]
                if stats['decreasing'] or stats['repeated']:
                    notes.append(issue('station_order_' + role, f"The mapped {role} contains {stats['decreasing']} decreasing and {stats['repeated']} repeated steps; confirm source ordering."))
            except (IndexError, ValueError, TypeError):
                pass
    dataset["issues"] = notes
    dataset["status"] = "retained/unresolved" if notes else "ready"
    dataset["family"] = family
    dataset["references"] = references
    dataset["mapping"] = mapping
    dataset["well"] = options.get("well") or dataset.get("metadata", {}).get("well", "")
    return dataset


def parse(path, output, family, options=None, companions=None):
    options = options or {}
    kind = detect(path, family)
    allowed={'las':{'logs','arrays'},'dlis':{'logs','arrays','checkshots'},'lis':{'logs','arrays'},'segy':{'seismic','checkshots','arrays'},'su':{'seismic','checkshots','arrays'},'zgy':{'seismic'},'sgz':{'seismic'},'vds':{'seismic'},'gocad':{'surfaces','models','spatial'},'surface':{'surfaces','models','spatial'},'resqml':{'models'},'eclipse':{'models'},'grdecl':{'models'},'pdf':{'documents'},'docx':{'documents'},'image':{'documents','core','arrays'}}
    if kind in allowed and family not in allowed[kind]:
        raise ImportProblem('data_type_mismatch','This file belongs to a different data family. Retain the original and import it in a separate matching data-type group.')
    if kind == "las":
        datasets = read_las(path, output, options)
    elif kind in ("table", "xlsx", "xls"):
        datasets = read_table(path, output, options)
    else:
        from ingestion_native import read_native
        datasets = read_native(kind, path, output, options, companions or {})
    for dataset in datasets:
        validate_mapping(dataset, family, options)
    return datasets
